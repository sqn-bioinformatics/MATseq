"""Assemble the manuscript's Table 2 and Supplementary Tables S1-S7."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

from .config import CONDITION_ORDER

DESEQ2_STATS = ["baseMean", "log2FoldChange", "lfcSE", "stat", "pvalue", "padj"]

S1_LIGAND_ORDER = [
    ("train_ligands", "Fla-PA", "Fla-PA"),
    ("train_ligands", "LPS", "LPS"),
    ("train_ligands", "PGN", "PGN"),
    ("train_ligands", "R848", "R848"),
    ("train_ligands", "Pam3", "Pam3"),
    ("additional_ligands", "LTA", "LTA"),
    ("additional_ligands", "MPLA", "MPLA"),
    ("additional_ligands", "Pam2", "Pam2"),
    ("bacterial_ligands", "HK E.coli", "HKEB"),
    ("bacterial_ligands", "HK S.aureus", "HKSA"),
    ("test_ligands", "Fla-PA", "test_Fla-PA"),
    ("test_ligands", "LPS", "test_LPS"),
    ("test_ligands", "Pam3", "test_Pam3"),
    ("test_ligands", "R848", "test_R848"),
]

S2_LIGAND_ORDER = [
    ("train_ligands", "Fla-PA", "Fla-PA"),
    ("train_ligands", "LPS", "LPS"),
    ("train_ligands", "PGN", "PGN"),
    ("train_ligands", "R848", "R848"),
    ("train_ligands", "Pam3", "Pam3"),
]

S4_CONDITIONS = [
    ("Intersect Feature Selection and Differentially Expressed Genes",
     "de_intersect_fs_go_terms.csv"),
    ("Feature Selection Genes not Differentially Expressed",
     "fs_only_go_terms.csv"),
]

# Reporter-assay source data (S5/S6): copied verbatim bar one renamed column.
REPORTER_TABLES = {
    "S5": ("SupplementaryTable5.csv",
           {"Concentration_(EU_mL)": "Concentration_LPS_(eq_mL)"}),
    "S6": ("SupplementaryTable6.csv",
           {"Concentration_(ng_mL)": "Concentration_Pam3_(ng_mL)"}),
}

CONDITION_LABELS = {
    "all_genes": "all genes",
    "feature_selection": "feature selection",
    "fs_plus_de": "feature selection + de",
    "random_selected": "randomly selected",
}
MODEL_ORDER = ["LinearSVC", "SGDClassifier", "LogisticRegression",
               "RandomForest", "XGBoost"]
MODEL_DISPLAY_NAMES = {
    "LinearSVC": "Linear SVC",
    "SGDClassifier": "SGD",
    "LogisticRegression": "LogisticRegression",
    "RandomForest": "Random Forest",
    "XGBoost": "XGBoost",
}
METRICS = [("accuracy", "Accuracy"), ("f1", "F1"),
           ("precision", "Precision"), ("recall", "Recall")]
TABLE2_COLUMNS = ["Condition", "Model", "Accuracy", "F1", "Precision", "Recall",
                  "Training time"]


def format_table2(raw_csv: Path, csv_path: Path,
                  xlsx_path: Path | None = None) -> pd.DataFrame:
    """Format a nested-CV summary into Table 2 (or its no-Fla-PA counterpart)."""
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    df = pd.read_csv(raw_csv)
    df["model"] = pd.Categorical(df["model"], MODEL_ORDER, ordered=True)
    df["condition"] = pd.Categorical(df["condition"], CONDITION_ORDER, ordered=True)
    df = df.sort_values(["condition", "model"]).reset_index(drop=True)

    fmt = pd.DataFrame([
        {
            "Condition": (f"{CONDITION_LABELS[r['condition']]} "
                          f"({int(round(r['n_genes_mean'])):,})"),
            "Model": MODEL_DISPLAY_NAMES[r["model"]],
            **{header: f"{r[f'{key}_mean']:.2f} ± {r[f'{key}_std']:.2f}"
               for key, header in METRICS},
            "Training time": f"{r['training_time_mean']:.0f} s",
        }
        for _, r in df.iterrows()
    ])[TABLE2_COLUMNS]
    fmt["Condition"] = fmt["Condition"].mask(fmt["Condition"].duplicated(), "")
    fmt.to_csv(csv_path, index=False)
    if xlsx_path is None:
        return fmt

    wb = Workbook()
    ws = wb.active
    ws.append(TABLE2_COLUMNS)
    for row in fmt.itertuples(index=False):
        ws.append(list(row))
    font = Font(name="Calibri", size=9, color="FF000000")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fills = [PatternFill("solid", fgColor="FFFFFFFF"),
             PatternFill("solid", fgColor=Color(theme=0, tint=-0.0499893185216834))]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(TABLE2_COLUMNS)):
        group = (row[0].row - 2) // len(MODEL_ORDER)
        for cell in row:
            cell.font = font
            cell.alignment = alignment
            cell.fill = fills[0] if row[0].row == 1 else fills[group % 2]
            if row[0].row == 1:
                cell.border = Border(bottom=Side(style="thin"))
    ws.row_dimensions[1].height = 24
    ws.column_dimensions["A"].width = 14.28515625
    for group in range(len(CONDITION_ORDER)):
        first = 2 + group * len(MODEL_ORDER)
        ws.merge_cells(start_row=first, start_column=1,
                       end_row=first + len(MODEL_ORDER) - 1, end_column=1)
    wb.save(xlsx_path)
    return fmt


def assemble_supplementary_table_1(de_dir: Path, output_dir: Path) -> pd.DataFrame:
    """Wide per-ligand DESeq2 statistics (S1)."""
    merged = None
    for subset, file_key, suffix in S1_LIGAND_ORDER:
        block = pd.read_csv(de_dir / subset / f"{file_key}_deseq2_results.csv")
        block = block.rename(columns={block.columns[0]: "gene"})
        block = block.rename(columns={c: f"{c}_{suffix}" for c in DESEQ2_STATS})
        merged = block if merged is None else merged.merge(block, on="gene", how="outer")
    merged.to_csv(output_dir / "SupplementaryTable1.csv", index=False)
    return merged


def assemble_supplementary_table_2(go_dir: Path, output_dir: Path) -> pd.DataFrame:
    """Per-ligand GO enrichment (S2)."""
    frames = []
    for subset, file_key, label in S2_LIGAND_ORDER:
        df = pd.read_csv(go_dir / subset / f"{file_key}_go_terms.csv")
        df.insert(0, "Ligand", label)
        frames.append(df)
    merged = pd.concat(frames, axis=0, ignore_index=True)
    merged = merged.sort_values("fdr", ascending=True).reset_index(drop=True)
    merged.to_csv(output_dir / "SupplementaryTable2.csv", index=False)
    return merged


def assemble_supplementary_table_3(fs_de_dir: Path, output_dir: Path) -> pd.DataFrame:
    """Selected genes ranked by ExtraTrees importance (S3)."""
    src = pd.read_csv(fs_de_dir / "selected_vs_de_overlap_table.csv")
    src = src.sort_values("rank").reset_index(drop=True)
    out = pd.DataFrame({
        "Rank": src["rank"].astype(int),
        "Gene": src["gene"],
        "Differentially_Expressed": src["in_de"].astype(bool),
    })
    out.to_csv(output_dir / "SupplementaryTable3.csv", index=False)
    return out


def assemble_supplementary_table_4(go_dir: Path, output_dir: Path) -> pd.DataFrame:
    """By-condition GO enrichment (S4)."""
    frames = []
    for label, fname in S4_CONDITIONS:
        df = pd.read_csv(go_dir / fname)
        df.insert(0, "Condition", label)
        frames.append(df)
    merged = pd.concat(frames, axis=0, ignore_index=True)
    merged.to_csv(output_dir / "SupplementaryTable4.csv", index=False)
    return merged


def assemble_supplementary_table_7(feature_selection_dir: Path,
                                   output_dir: Path) -> pd.DataFrame:
    """Forest/KMeans ARI scan (S7)."""
    src = pd.read_csv(feature_selection_dir / "forest_kmeans.csv")
    src = src.sort_values("n_selected").reset_index(drop=True)
    out = pd.DataFrame({"n_selected_genes": src["n_selected"]})
    for raw_col in src.columns[src.columns.str.startswith("ari_seed_")]:
        out[raw_col.replace("ari_seed_", "ARI_seed_")] = src[raw_col]
    out["ARI_mean"] = src["ari_mean"]
    out["ARI_std"] = src["ari_std"]
    out.to_csv(output_dir / "SupplementaryTable7.csv", index=False)
    return out


def assemble_supplementary_tables(
    de_dir: Path, go_dir: Path, fs_de_dir: Path,
    feature_selection_dir: Path, supp_data_dir: Path, output_dir: Path,
) -> dict[str, pd.DataFrame]:
    """Regenerate Supplementary Tables S1-S7 from raw pipeline outputs."""
    output_dir.mkdir(parents=True, exist_ok=True)
    tables = {
        "S1": assemble_supplementary_table_1(de_dir, output_dir),
        "S2": assemble_supplementary_table_2(go_dir, output_dir),
        "S3": assemble_supplementary_table_3(fs_de_dir, output_dir),
        "S4": assemble_supplementary_table_4(go_dir, output_dir),
    }
    for key, (filename, renames) in REPORTER_TABLES.items():
        tables[key] = pd.read_csv(supp_data_dir / filename).rename(columns=renames)
        tables[key].to_csv(output_dir / filename, index=False)
    tables["S7"] = assemble_supplementary_table_7(feature_selection_dir, output_dir)
    return tables
